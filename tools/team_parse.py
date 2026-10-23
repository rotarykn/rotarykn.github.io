import openpyxl, re, json, sys
def clean(s): return re.sub(r'[\s　]+', '', str(s or ''))
def nm(s):
    s = clean(s); return '吳　狄' if s == '吳狄' else s
def role(s):
    m = re.fullmatch(r'[\(（](.+)[\)）]', clean(s)); return m.group(1) if m else None
def gname(s):
    s = clean(s).lstrip('˙・·'); return '青少年服務委員會' if s == '青少年服務主委' else s
def names(txt, word):
    out = []
    for line in re.split(r'[\n、,，]', txt):
        line = line.replace(word, '').replace('：', '').replace(':', '')
        m = re.search(r'[一-鿿](?:[ 　]*[一-鿿]){1,3}', line)
        if m: out.append(nm(m.group(0)))
    return out

def parse(path, no, year):
    ws = openpyxl.load_workbook(path, data_only=True).active
    V = lambda c, r: ws[f'{c}{r}'].value
    pres, dirs, sups, hdr = None, [], [], None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None: continue
            s = str(cell.value)
            if clean(s) == '社長': pres = nm(ws.cell(cell.row, cell.column + 1).value)
            if clean(s) == '行政管理委員會': hdr = cell.row
            if '理事' in s and clean(s) != '理事' and not dirs: dirs = names(s[s.find('理事'):].split('監事')[0], '理事')
            if '監事' in s and clean(s) != '監事' and not sups: sups = names(s[s.find('監事'):], '監事')
    cols = [('B', 'C', 'D'), ('E', 'F', 'G'), ('H', 'I', 'J'), ('K', 'L', 'M'), ('N', 'O', None)]
    cmts = []
    for gi, (gc, nc, rc) in enumerate(cols):
        cm = {'n': gname(V(gc, hdr)), 'ch': nm(V(nc, hdr + 1)), 'c': []}
        wrap = top = cur = None
        for r in range(hdr + 2, ws.max_row + 1):
            g, n = V(gc, r), V(nc, r)
            x = V(rc, r) if rc else None
            if clean(g):
                grp = {'n': gname(g), 'p': [], 'c': []}
                if gi == 0 and clean(g) == '社務服務委員會':
                    wrap = grp; cm['c'].append(grp); cur = None; continue
                if wrap is not None: wrap['c'].append(grp)
                elif str(g).strip()[0] in '˙・·' and top is not None: top['c'].append(grp)
                else: cm['c'].append(grp); top = grp
                cur = grp
            if cur is None: continue
            rl = role(x) if clean(x) else None
            if clean(n): cur['p'].append([nm(n), rl] if rl else [nm(n)])
            elif rl and rl != '顧問' and cur['p']: pass
            if clean(x) and not rl and gi == 2: cur['p'].append([nm(x)])
        cmts.append(cm)
    def tidy(g):
        if not g.get('p'): g.pop('p', None)
        if not g.get('c'): g.pop('c', None)
        for c in g.get('c', []): tidy(c)
    for c in cmts:
        for g in c['c']: tidy(g)
    return {'no': no, 'year': year, 'pres': pres, 'dirs': dirs, 'sups': sups, 'cmts': cmts}

if __name__ == '__main__':
    out = [parse(p, int(n), y) for p, n, y in zip(sys.argv[1::3], sys.argv[2::3], sys.argv[3::3])]
    print(json.dumps(out, ensure_ascii=False))
